from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from django.http import HttpResponse
from openpyxl.utils import get_column_letter
from academics.models import Student, Promotion, UE
from evaluations.models import Grade
from io import BytesIO


def export_deliberation_excel(request, promotion_id):

    wb = Workbook()
    ws = wb.active
    ws.title = "Deliberation"

    # ======================
    # 🎨 STYLES
    # ======================
    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    fill_ue = PatternFill(start_color="CCCCCC", fill_type="solid")
    fill_course = PatternFill(start_color="EEEEEE", fill_type="solid")

    # ======================
    # 🏫 HEADER
    # ======================
    promotion = Promotion.objects.get(id=promotion_id)

    ws.merge_cells('A1:Z1')
    ws['A1'] = "HAUTE ECOLE DE COMMERCE DE KINSHASA"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center

    ws.merge_cells('A2:Z2')
    ws['A2'] = "GRILLE DE DELIBERATION"
    ws['A2'].alignment = center

    ws.merge_cells('A3:Z3')
    ws['A3'] = f"CLASSE : {promotion.name}"
    ws['A3'].alignment = center

    # ======================
    # 📊 STRUCTURE FIXE
    # ======================
    students = Student.objects.filter(promotion=promotion)

    ues = UE.objects.prefetch_related('course_set').all()

    row_ue = 5
    row_course = 6
    row_sub = 7

    # COLONNES FIXES
    headers = ["N°", "POSTNOM", "PRENOM"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row_sub, column=i).value = h

    col = 4

    # ======================
    # 🧱 CONSTRUCTION TABLE
    # ======================
    for ue in ues:
        courses = ue.course_set.all()
        start_col = col

        for course in courses:

            # Fusion nom cours
            ws.merge_cells(start_row=row_course, start_column=col,
                           end_row=row_course, end_column=col+2)

            c = ws.cell(row=row_course, column=col)
            c.value = course.name
            c.fill = fill_course
            c.font = bold
            c.alignment = center

            # Sous colonnes
            ws.cell(row=row_sub, column=col).value = "Note"
            ws.cell(row=row_sub, column=col+1).value = "Cr"
            ws.cell(row=row_sub, column=col+2).value = "TNP"

            col += 3

        # Ajouter MOY UE
        ws.cell(row=row_sub, column=col).value = "MOY UE"
        col += 1

        end_col = col - 1

        # Fusion UE
        ws.merge_cells(start_row=row_ue, start_column=start_col,
                       end_row=row_ue, end_column=end_col)

        u = ws.cell(row=row_ue, column=start_col)
        u.value = ue.name
        u.fill = fill_ue
        u.font = bold
        u.alignment = center

    # Colonnes finales
    ws.cell(row=row_sub, column=col).value = "MOY GEN"
    ws.cell(row=row_sub, column=col+1).value = "DECISION"

    col_moy = col

    # ======================
    # 👨🎓 DATA
    # ======================
    row = 8
    index = 1

    for student in students:

        ws.cell(row=row, column=1).value = index
        ws.cell(row=row, column=2).value = student.postnom
        ws.cell(row=row, column=3).value = student.prenom

        col = 4
        total_tnp = 0
        total_credit = 0

        for ue in ues:
            courses = ue.course_set.all()

            ue_tnp = 0
            ue_credit = 0

            for course in courses:
                grade = Grade.objects.filter(student=student, course=course).first()

                note = grade.note_finale if grade and grade.note_finale else 0
                credit = course.credit
                tnp = note * credit

                ws.cell(row=row, column=col).value = note
                ws.cell(row=row, column=col+1).value = credit
                ws.cell(row=row, column=col+2).value = tnp

                ue_tnp += tnp
                ue_credit += credit

                col += 3

            # MOY UE
            moy_ue = ue_tnp / ue_credit if ue_credit else 0
            ws.cell(row=row, column=col).value = round(moy_ue, 2)

            col += 1

            total_tnp += ue_tnp
            total_credit += ue_credit

        # MOY GENERALE
        moy = total_tnp / total_credit if total_credit else 0

        ws.cell(row=row, column=col_moy).value = round(moy, 2)

        decision = "ADM" if moy >= 10 else "AJ" if moy >= 8 else "DEF"
        ws.cell(row=row, column=col_moy+1).value = decision

        row += 1
        index += 1

    # ======================
    # 🎨 STYLE GLOBAL
    # ======================
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.border = border
            cell.alignment = center

    # LARGEUR AUTO
    for i, col_cells in enumerate(ws.columns, 1):
        max_length = 0
        for cell in col_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(i)].width = max_length + 2

    # ======================
    # 📥 EXPORT
    # ======================
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = f'attachment; filename=DELIBERATION_{promotion.name}.xlsx'

    return response